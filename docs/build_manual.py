"""Generate the end-user Word manual with the bundled python-docx runtime."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'artifacts' / 'ScreenWatch使用说明书.docx'
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(.68)
section.bottom_margin = Inches(.66)
section.left_margin = Inches(.78)
section.right_margin = Inches(.78)
section.footer_distance = Inches(.3)


def font(style, size, bold=False):
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().rFonts
    for name in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        fonts.attrib.pop(qn('w:' + name), None)
    fonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    for border in style.element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)


for name, size, bold in [('Normal', 11, False), ('Title', 25, True),
                         ('Heading 1', 17, True), ('Heading 2', 12, True), ('Footer', 9, False)]:
    font(doc.styles[name], size, bold)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = Pt(17)
normal.space_after = Pt(7)
for name in ['Heading 1', 'Heading 2']:
    pf = doc.styles[name].paragraph_format
    pf.keep_with_next = True
    pf.space_before = Pt(12)
    pf.space_after = Pt(7)
    pf.line_spacing = Pt(23 if name == 'Heading 1' else 18)
doc.styles['Title'].paragraph_format.space_after = Pt(8)
doc.styles['Title'].paragraph_format.line_spacing = Pt(32)


def p(text='', *, bold_lead=None, size=None, after=None):
    paragraph = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        paragraph.add_run(bold_lead).bold = True
        paragraph.add_run(text[len(bold_lead):])
    else:
        paragraph.add_run(text)
    if size:
        for run in paragraph.runs:
            run.font.size = Pt(size)
    if after is not None:
        paragraph.paragraph_format.space_after = Pt(after)
    return paragraph


def heading(text, level=1):
    return doc.add_heading(text, level)


def step(number, title, body):
    para = p(f'{number}  {title}　{body}', bold_lead=f'{number}  {title}')
    para.paragraph_format.keep_together = True
    return para


def new_page(title):
    h = heading(title)
    h.paragraph_format.page_break_before = True
    h.paragraph_format.space_before = Pt(0)


def table(headers, rows, widths):
    result = doc.add_table(rows=1, cols=len(headers))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    for column, width in zip(result.columns, widths):
        column.width = Inches(width)
    props = result._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        edge = OxmlElement('w:' + side)
        for key, value in [('val', 'single'), ('sz', '5'), ('color', 'D9D9D9')]:
            edge.set(qn('w:' + key), value)
        borders.append(edge)
    props.append(borders)
    repeat = OxmlElement('w:tblHeader')
    result.rows[0]._tr.get_or_add_trPr().append(repeat)
    for index, values in enumerate([headers] + rows):
        row = result.rows[0] if index == 0 else result.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for col, (cell, text, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'E7EDF5' if index == 0 else 'FFFFFF')
            tcpr.append(shade)
            margins = OxmlElement('w:tcMar')
            for side, value in [('top', '95'), ('bottom', '95'), ('left', '110'), ('right', '110')]:
                element = OxmlElement('w:' + side)
                element.set(qn('w:w'), value)
                element.set(qn('w:type'), 'dxa')
                margins.append(element)
            tcpr.append(margins)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = Pt(15.5)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER if col == 1 and len(headers) == 3 else WD_ALIGN_PARAGRAPH.LEFT
            run = para.add_run(text)
            run.font.size = Pt(10.5)
            run.bold = index == 0
    p('', after=1).paragraph_format.line_spacing = .3
    return result


doc.core_properties.title = 'ScreenWatch 使用说明书'
doc.core_properties.subject = 'Windows 通用画面监控程序的安装配置与日常使用'
doc.core_properties.author = 'ScreenWatch'
doc.core_properties.keywords = 'ScreenWatch Windows 使用说明 画面监控'

doc.add_paragraph('ScreenWatch 使用说明书', 'Title')
p('软件版本 1.0.0    适用平台 Windows x64    更新日期 2026年9月17日', size=9.5, after=13)
p('本手册介绍如何监控屏幕上的固定区域：保存一张参考画面，在画面再次出现时收到本地提醒。按照下列步骤完成配置后，即可将程序隐藏到系统托盘运行。')
p('使用前请确认：被监控的 Chrome 或其他应用必须保持可见，目标区域不能被遮挡。程序隐藏到托盘后可以继续监控；目标窗口最小化或电脑锁屏时无法继续读取原画面。', bold_lead='使用前请确认：')

heading('1 安装与首次使用')
step(1, '解压并启动', '将 ScreenWatch-1.0.0-win-x64.zip 复制到 Windows 电脑，完整解压后双击 ScreenWatch.exe。不要直接在压缩包内运行。')
p('适用于受 .NET 10 支持的 Windows 10/11 x64 版本，无需另装 .NET、Python、Java 或浏览器驱动，无需管理员权限。首次启动会解压内置组件，请稍候。', size=10.5)
step(2, '准备目标画面', '打开要监控的应用或网页，固定窗口位置、页面滚动位置和缩放比例。首次使用建议先按第 3 节完成离线演示。')
step(3, '框选监控区域', '点击“框选区域”，主窗口暂时隐藏。在目标上按住鼠标左键拖动，松开确认；按 Esc 或右键取消。区域尽量贴合目标，并限制在同一块显示器内。')
step(4, '保存参考画面', '让希望识别的画面显示在框选区域内，点击“保存当前画面为参考”。确认“参考画面”预览正确；也可点击“导入参考图片”，选择与区域像素尺寸完全一致的 PNG、JPEG 或 BMP。')
step(5, '测试匹配', '点击“测试一次匹配”，检查相似度和“最近取样”。分别测试目标画面和其他画面：目标应达到阈值，其他画面应低于阈值。手动测试不会触发提醒。')
step(6, '开始监控', '首次可保留默认参数，点击“开始监控”。主窗口会隐藏到托盘。连续命中达到设定次数后，程序弹出提醒，并根据选项播放提示音、保存截图。')
p('重新框选区域会使原参考图片失效，必须重新保存或导入参考图片，才能再次开始监控。', bold_lead='重新框选区域')

new_page('2 参数设置与日常操作')
p('建议先使用默认设置。若出现误提醒或漏提醒，先缩小监控区域、去掉动画和时间等变化内容，再通过“测试一次匹配”调整阈值。相似度是画面比较分数，不是识别准确率。')
table(['设置', '默认值', '说明'], [
    ['间隔（毫秒）', '1000', '可设 250 至 60000。数值越小，取样越频繁；上一轮未结束时会跳过，不堆积任务。'],
    ['阈值（%）', '95', '可设 50 至 100。越高越严格；降低阈值可能增加误匹配。'],
    ['连续命中次数', '2', '可设 1 至 20。连续多次匹配后才提醒，用于过滤短暂闪烁。'],
    ['消失确认次数', '2', '可设 1 至 20。连续多次不匹配后，才允许识别下一次出现。'],
    ['冷却（秒）', '30', '可设 0 至 3600。限制两次提醒的最短间隔。'],
    ['提示音', '开启', '命中时播放系统提示音；受系统音量设置影响。'],
    ['命中时保存截图', '开启', '命中时保存监控区域的 PNG 图片，便于事后核对。'],
], [1.4, .72, 4.82])
heading('一次出现只提醒一次', 2)
p('画面持续匹配时，即使冷却结束，也不会重复提醒。只有画面连续消失达到设定次数后，再次出现才可能触发下一次提醒。')
p('如果新的出现发生在冷却期内，程序会等到冷却结束且画面仍然匹配时再提醒；如果提前消失，则不会补发。每次手动“开始监控”都会重置命中状态和冷却计时。')
heading('窗口与托盘操作', 2)
table(['操作', '结果'], [
    ['关闭或最小化主窗口', '隐藏到托盘，不等于退出程序。'],
    ['双击托盘图标', '打开主窗口并停止监控，避免主窗口遮挡目标；再次点击“开始监控”恢复。'],
    ['右键托盘图标', '可打开主窗口、开始或停止监控、打开数据目录、退出程序。'],
    ['打开“数据目录”', '若正在监控，会先停止；检查完成后需要手动重新开始。'],
    ['点击“退出”', '保存设置并真正结束程序。重新启动后恢复配置，但不会自动开始监控。'],
], [1.65, 5.29])

new_page('3 离线演示与常见问题')
heading('用演示页面验证流程', 2)
p('演示页面只在本机运行，无需联网。将 ScreenWatch 和演示页面配合使用，可以先确认框选、匹配与提醒流程。')
step(1, '打开演示页', '用 Chrome 打开解压目录中的 demo/index.html，点击“显示已完成”。')
step(2, '建立参考', '在 ScreenWatch 中框选整张状态卡，包括边框，然后保存当前画面为参考。手动测试时，相似度应接近 100%。')
step(3, '设置测试参数', '间隔设为 1000 毫秒、阈值 95%、连续命中 2 次、消失确认 2 次、冷却 3 秒。')
step(4, '观察自动提醒', '在演示页点击“自动切换”，再到 ScreenWatch 点击“开始监控”。让演示页保持可见，状态每 6 秒切换一次；“已完成”稳定出现后应提醒一次。')
step(5, '验证不会重复', '让“已完成”保持不变，等待超过冷却时间，应不再提醒。切换为“等待中”并持续至少两次取样，再切回“已完成”，应允许新的提醒。')
heading('常见问题处理', 2)
p('找不到主窗口　查看 Windows 右下角托盘及隐藏图标区域，双击 ScreenWatch 图标。若再次启动时提示“已在运行”，原实例通常仍在托盘中。', bold_lead='找不到主窗口')
p('“开始监控”不可点击　确认已框选区域并保存或导入参考图片。重新框选后需要重新设置参考；正在框选、导入或取样时，请先完成当前操作。', bold_lead='“开始监控”不可点击')
p('目标出现却没有提醒　先查看“测试一次匹配”的分数；检查区域是否偏移、缩放是否变化、连续命中次数是否满足，以及是否仍处于冷却期。查看日志确认是否已命中。', bold_lead='目标出现却没有提醒')
p('日志已命中但没有气泡或声音　检查 Windows 通知权限、“请勿打扰”、系统音量和程序的“提示音”选项。系统可能隐藏通知气泡，但日志仍会记录命中。', bold_lead='日志已命中但没有气泡或声音')
p('出现误提醒　先缩小区域，避开大块空白、动态时间、动画和通知区域，再提高阈值或增加连续命中次数。用目标与非目标画面分别测试。', bold_lead='出现误提醒')
p('参考图片导入失败　确认图片未损坏、文件不超过 64 MB，且像素宽高与框选区域完全一致。最直接的方法是使用“保存当前画面为参考”重新取样。', bold_lead='参考图片导入失败')

new_page('4 运行条件与数据管理')
heading('保持监控稳定', 2)
p('程序记录的是固定屏幕坐标当前显示的内容，不会跟随窗口移动。移动窗口、滚动页面、切换标签、调整浏览器或 Windows 缩放后，请停止监控并检查区域；必要时重新框选和保存参考。')
p('锁屏、会话断开、休眠或恢复、显示器布局改变，会停止监控。回到桌面后先确认目标画面正确，再手动“开始监控”。截图发生错误时也会停止，并提示原因。')
p('不要让其他窗口、通知气泡或浮层遮挡目标。建议避开屏幕右下角的通知区域。程序每次只监控一个区域和一张参考图片，不进行整屏搜索或文字识别。')
heading('查找配置与截图', 2)
p('点击“数据目录”可打开本地数据文件夹，也可在 Windows 文件资源管理器地址栏中输入以下路径：')
p('%LOCALAPPDATA%\\ScreenWatch\\', size=11)
table(['文件或目录', '用途'], [
    ['settings.json', '监控区域、参考图片文件名及提醒参数。'],
    ['reference-<随机编号>.png', '当前参考图片。'],
    [r'captures\match-*.png', '命中时保存的区域截图。'],
    [r'logs\YYYY-MM-DD.log', '开始、停止、匹配提醒及错误记录。'],
], [2.5, 4.44])
p('数据只保存在本机，程序没有云上传功能。命中截图保留最近 7 天，最多 200 张且总量不超过 256 MiB；程序启动及保存命中截图时清理。日志保留 30 天，每天最多约 4 MiB，超出后轮换。')
heading('备份 恢复与卸载', 2)
p('备份　先退出程序，再复制整个 ScreenWatch 数据目录。迁移到另一台电脑后，即使恢复了设置，也应重新检查区域，因为显示器位置、分辨率和缩放可能不同。', bold_lead='备份')
p('配置损坏　程序会报错并退出，不会自动覆盖旧配置。先备份数据目录，再将 settings.json 移到其他位置，重启后重新配置。', bold_lead='配置损坏')
p('卸载　通过托盘菜单退出程序后删除解压目录即可；如不再需要历史记录，可另行删除上述数据目录。', bold_lead='卸载')
p('使用前验证　先用第 3 节演示流程确认当前电脑上的截图、通知和托盘操作正常，再开始日常监控。', bold_lead='使用前验证')

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.paragraph_format.space_after = Pt(0)
for text, field in [('ScreenWatch 使用说明书    第 ', None), ('', 'PAGE'), (' 页 / 共 ', None), ('', 'NUMPAGES'), (' 页', None)]:
    run = footer.add_run(text)
    run.font.size = Pt(9)
    if field:
        element = OxmlElement('w:fldSimple')
        element.set(qn('w:instr'), field)
        run._r.addnext(element)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
